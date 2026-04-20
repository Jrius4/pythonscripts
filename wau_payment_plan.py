#!/usr/bin/env python3
from __future__ import annotations

import argparse
import math
import random
import re
import unicodedata
import warnings
from dataclasses import dataclass
from datetime import datetime
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook


WAU_SHEET_NAME = "Data"
BENCHMARK_REQUIRED_COLUMNS = (
    "payment_id",
    "individual_id",
    "collector_name",
    "fsp_name",
    "entitlement_quantity",
    "delivered_quantity",
    "delivery_date",
    "reason_for_unsuccessful_payment",
    "admin1",
    "admin2",
    "admin3",
    "full_name",
    "given_name",
    "middle_name",
    "family_name",
    "sex",
    "birth_date",
    "phone_no",
    "birth_certificate_no",
    "national_id_no",
    "ss_hw_work_id_i_f",
    "ss_hw_lot_num_i_f",
    "ss_health_facility_name_i_f",
    "ss_hw_title_i_f",
    "ss_hw_hope_july_ref_i_f",
)


class DataQualityError(RuntimeError):
    """Raised when source workbooks do not support a reliable payment-plan output."""


@dataclass(frozen=True)
class AdminInfo:
    state: str
    state_pcode: str
    county: str
    county_pcode: str
    payam: str
    payam_pcode: str


def parse_args() -> argparse.Namespace:
    base_dir = Path(__file__).resolve().parent
    input_dir = base_dir / "inputs"

    parser = argparse.ArgumentParser(
        description=(
            "Generate a Wau payment plan from the Wau health-worker roster, "
            "benchmarking the layout and role-based amounts from an existing "
            "payment plan workbook."
        )
    )
    parser.add_argument(
        "--wau",
        type=Path,
        default=input_dir / "Wau.xlsx",
        help="Path to the Wau health-worker workbook.",
    )
    parser.add_argument(
        "--benchmark",
        type=Path,
        default=input_dir
        / "payment_plan_payment_list_PP-4040-26-00000280_FSP_Co-operative Bank of South Sudan Ltd_Cash on Site.xlsx",
        help="Path to the benchmark payment-plan workbook.",
    )
    parser.add_argument(
        "--orgunit",
        type=Path,
        default=input_dir / "OrgUnit_pcode.xlsx",
        help="Path to the OrgUnit workbook with State/County/Payam PCodes.",
    )
    parser.add_argument(
        "--output-dir",
        type=Path,
        default=input_dir,
        help="Directory for the generated payment-plan workbook.",
    )
    parser.add_argument(
        "--skip-count",
        type=int,
        default=10,
        help="How many eligible workers to intentionally omit before extras are added.",
    )
    parser.add_argument(
        "--extra-count",
        type=int,
        default=20,
        help="How many extra workers to add back from the remaining roster.",
    )
    parser.add_argument(
        "--seed",
        type=int,
        default=404026,
        help="Deterministic selection seed for omitted and extra workers.",
    )
    return parser.parse_args()


def is_missing(value: object) -> bool:
    return value is None or (isinstance(value, float) and math.isnan(value)) or pd.isna(value)


def clean_text(value: object) -> str | None:
    if is_missing(value):
        return None

    text = unicodedata.normalize("NFKC", str(value))
    text = text.replace("\xa0", " ")
    text = re.sub(r"\s+", " ", text).strip()
    return text or None


def canonical_admin_name(value: object) -> str | None:
    text = clean_text(value)
    if not text:
        return None

    text = text.casefold()
    text = re.sub(r"[^\w\s/-]", "", text)
    for suffix in (" state", " county", " payam"):
        if text.endswith(suffix):
            text = text[: -len(suffix)].strip()
    text = re.sub(r"\s+", " ", text).strip()
    return text or None


def display_admin_name(value: object) -> str | None:
    text = clean_text(value)
    if not text:
        return None

    text = re.sub(r"\s+(State|County|Payam)$", "", text, flags=re.IGNORECASE).strip()
    return text or None


def canonical_worker_key(individual_id: object, work_id: object, full_name: object) -> str:
    for candidate in (clean_text(individual_id), clean_text(work_id), clean_text(full_name)):
        if candidate:
            return candidate.casefold()
    raise DataQualityError("Encountered a Wau worker row without any usable identifier.")


def parse_bool(value: object) -> bool:
    if isinstance(value, bool):
        return value
    text = clean_text(value)
    if not text:
        return False
    return text.casefold() in {"true", "1", "yes", "y"}


def pick_preferred_value(row: pd.Series, base_column: str, updated_column: str | None = None) -> str | None:
    if updated_column and updated_column in row.index:
        updated = clean_text(row.get(updated_column))
        if updated:
            return updated
    return clean_text(row.get(base_column))


def format_birth_date(value: object) -> str | None:
    if is_missing(value):
        return None
    if isinstance(value, datetime):
        return value.strftime("%Y-%m-%d")

    parsed = pd.to_datetime(value, errors="coerce")
    if pd.isna(parsed):
        text = clean_text(value)
        return text
    return parsed.strftime("%Y-%m-%d")


def split_name(full_name: str | None) -> tuple[str | None, str | None, str | None]:
    name = clean_text(full_name)
    if not name:
        return None, None, None

    parts = name.split(" ")
    given_name = parts[0]
    family_name = parts[1] if len(parts) >= 2 else None
    middle_name = " ".join(parts[2:]) if len(parts) >= 3 else None
    return given_name, middle_name, family_name


def load_benchmark_metadata(path: Path) -> tuple[pd.DataFrame, str, list[str]]:
    warnings.filterwarnings(
        "ignore",
        message="Slicer List extension is not supported and will be removed",
    )
    benchmark = pd.read_excel(path, dtype=object)
    missing = [column for column in BENCHMARK_REQUIRED_COLUMNS if column not in benchmark.columns]
    if missing:
        raise DataQualityError(
            "Benchmark workbook is missing required columns: " + ", ".join(missing)
        )

    fsp_names = [name for name in benchmark["fsp_name"].dropna().map(clean_text).unique().tolist() if name]
    if not fsp_names:
        raise DataQualityError("Benchmark workbook does not contain any fsp_name values.")

    return benchmark, fsp_names[0], list(benchmark.columns)


def build_entitlement_map(benchmark: pd.DataFrame) -> dict[str, int]:
    working = benchmark[["ss_hw_title_i_f", "entitlement_quantity"]].copy()
    working["ss_hw_title_i_f"] = working["ss_hw_title_i_f"].map(clean_text)
    working = working.dropna(subset=["ss_hw_title_i_f", "entitlement_quantity"])
    working["entitlement_quantity"] = working["entitlement_quantity"].astype(float).astype(int)

    entitlement_map: dict[str, int] = {}
    for title, group in working.groupby("ss_hw_title_i_f", dropna=False):
        counts = (
            group["entitlement_quantity"]
            .value_counts()
            .rename_axis("amount")
            .reset_index(name="count")
            .sort_values(["count", "amount"], ascending=[False, False])
        )
        entitlement_map[title] = int(counts.iloc[0]["amount"])

    return entitlement_map


def build_admin_lookup(path: Path) -> dict[tuple[str, str, str], AdminInfo]:
    orgunit = pd.read_excel(path, dtype=object)
    required = ["State", "State_PCode", "County", "County_PCode", "Payam", "Payam_PCode"]
    missing = [column for column in required if column not in orgunit.columns]
    if missing:
        raise DataQualityError("OrgUnit workbook is missing required columns: " + ", ".join(missing))

    lookup: dict[tuple[str, str, str], AdminInfo] = {}
    for row in orgunit[required].itertuples(index=False):
        state, state_pcode, county, county_pcode, payam, payam_pcode = row
        key = (
            canonical_admin_name(state),
            canonical_admin_name(county),
            canonical_admin_name(payam),
        )
        if any(part is None for part in key):
            continue

        info = AdminInfo(
            state=clean_text(state) or "",
            state_pcode=clean_text(state_pcode) or "",
            county=clean_text(county) or "",
            county_pcode=clean_text(county_pcode) or "",
            payam=clean_text(payam) or "",
            payam_pcode=clean_text(payam_pcode) or "",
        )

        existing = lookup.get(key)
        if existing and existing != info:
            raise DataQualityError(
                "Conflicting admin mapping found in OrgUnit workbook for "
                f"{key!r}: {existing} vs {info}"
            )
        lookup[key] = info

    if not lookup:
        raise DataQualityError("No usable State/County/Payam mappings found in OrgUnit workbook.")

    return lookup


def prepare_wau_roster(path: Path) -> pd.DataFrame:
    warnings.filterwarnings(
        "ignore",
        message="Slicer List extension is not supported and will be removed",
    )
    wau = pd.read_excel(path, sheet_name=WAU_SHEET_NAME, dtype=object)

    records: list[dict[str, object]] = []
    for _, row in wau.iterrows():
        full_name = clean_text(row.get("Collector Name"))
        individual_id = clean_text(row.get("unicef_id"))
        work_id = clean_text(row.get("HH ID"))

        record = {
            "worker_key": canonical_worker_key(individual_id, work_id, full_name),
            "individual_id": individual_id,
            "work_id": work_id,
            "collector_name": full_name,
            "full_name": full_name,
            "sex": pick_preferred_value(row, "Sex"),
            "birth_date": format_birth_date(row.get("Birth Date")),
            "phone_no": pick_preferred_value(row, "Phone Number", "Phone Number UPD"),
            "health_facility": pick_preferred_value(row, "Health Facility", "Health Facility UPD"),
            "title": pick_preferred_value(row, "Title", "Title UPD"),
            "state": pick_preferred_value(row, "State", "State UPD"),
            "county": pick_preferred_value(row, "County", "County UPD"),
            "payam": pick_preferred_value(row, "Payam", "Payam UPD"),
            "lot": pick_preferred_value(row, "Lot", "Lot UPD"),
            "eligible": parse_bool(row.get("ELIGIBLE IN CYCLE 7 (JULY)")),
            "hope_reference": individual_id,
        }
        records.append(record)

    roster = pd.DataFrame(records)
    roster["_completeness"] = roster.notna().sum(axis=1)
    roster = roster.sort_values(
        ["_completeness", "eligible", "individual_id", "collector_name"],
        ascending=[False, False, True, True],
        na_position="last",
    )
    roster = roster.drop_duplicates(subset=["worker_key"], keep="first").copy()
    roster = roster.drop(columns=["_completeness"]).reset_index(drop=True)

    required = ["individual_id", "collector_name", "title", "state", "county", "payam", "lot"]
    missing_rows = roster[roster[required].isna().any(axis=1)]
    if not missing_rows.empty:
        raise DataQualityError(
            "Some deduplicated Wau workers are missing required fields. "
            f"First affected worker: {missing_rows.iloc[0].to_dict()}"
        )

    return roster


def select_workers(roster: pd.DataFrame, skip_count: int, extra_count: int, seed: int) -> tuple[pd.DataFrame, dict[str, int]]:
    rng = random.Random(seed)

    eligible = roster[roster["eligible"]].copy()
    ineligible = roster[~roster["eligible"]].copy()

    eligible_indices = eligible.index.tolist()
    skipped_indices = sorted(rng.sample(eligible_indices, min(skip_count, len(eligible_indices))))

    skipped = eligible.loc[skipped_indices].copy()
    selected = eligible.drop(index=skipped_indices).copy()

    extra_frames: list[pd.DataFrame] = []
    extra_added = 0

    for pool in (ineligible, skipped):
        if extra_added >= extra_count or pool.empty:
            continue

        pool_indices = pool.index.tolist()
        chosen_indices = sorted(
            rng.sample(pool_indices, min(extra_count - extra_added, len(pool_indices)))
        )
        extra_frames.append(pool.loc[chosen_indices].copy())
        extra_added += len(chosen_indices)

    extras = pd.concat(extra_frames, ignore_index=False) if extra_frames else roster.iloc[0:0].copy()

    final = pd.concat([selected, extras], ignore_index=False)
    final = final.drop_duplicates(subset=["worker_key"], keep="first").copy()
    final = final.reset_index(drop=True)

    stats = {
        "deduplicated_workers": int(len(roster)),
        "eligible_workers": int(len(eligible)),
        "ineligible_workers": int(len(ineligible)),
        "skipped_workers": int(len(skipped)),
        "extra_workers_added": int(len(extras)),
        "selected_workers": int(len(final)),
    }
    return final, stats


def build_output_rows(
    workers: pd.DataFrame,
    entitlement_map: dict[str, int],
    admin_lookup: dict[tuple[str, str, str], AdminInfo],
    fsp_name: str,
    run_stamp: str,
) -> list[dict[str, object]]:
    rows: list[dict[str, object]] = []

    for row in workers.itertuples(index=False):
        title = clean_text(row.title)
        if title not in entitlement_map:
            raise DataQualityError(f"No entitlement benchmark found for title: {title}")

        admin_key = (
            canonical_admin_name(row.state),
            canonical_admin_name(row.county),
            canonical_admin_name(row.payam),
        )
        admin = admin_lookup.get(admin_key)
        if admin is None:
            raise DataQualityError(
                "No OrgUnit admin mapping found for "
                f"State={row.state!r}, County={row.county!r}, Payam={row.payam!r}"
            )

        given_name, middle_name, family_name = split_name(row.full_name)

        rows.append(
            {
                "payment_id": None,
                "individual_id": clean_text(row.individual_id),
                "collector_name": clean_text(row.collector_name),
                "fsp_name": fsp_name,
                "entitlement_quantity": int(entitlement_map[title]),
                "delivered_quantity": None,
                "delivery_date": None,
                "reason_for_unsuccessful_payment": None,
                "admin1": f"{admin.state_pcode} - {display_admin_name(admin.state)}",
                "admin2": f"{admin.county_pcode} - {display_admin_name(admin.county)}",
                "admin3": f"{admin.payam_pcode} - {display_admin_name(admin.payam)}",
                "full_name": clean_text(row.full_name),
                "given_name": given_name,
                "middle_name": middle_name,
                "family_name": family_name,
                "sex": clean_text(row.sex),
                "birth_date": clean_text(row.birth_date),
                "phone_no": clean_text(row.phone_no),
                "birth_certificate_no": None,
                "national_id_no": None,
                "ss_hw_work_id_i_f": clean_text(row.work_id),
                "ss_hw_lot_num_i_f": clean_text(row.lot),
                "ss_health_facility_name_i_f": clean_text(row.health_facility),
                "ss_hw_title_i_f": title,
                "ss_hw_hope_july_ref_i_f": clean_text(row.hope_reference),
            }
        )

    rows.sort(
        key=lambda item: (
            item["admin2"] or "",
            item["admin3"] or "",
            item["ss_health_facility_name_i_f"] or "",
            item["ss_hw_title_i_f"] or "",
            item["full_name"] or "",
        )
    )
    for sequence, item in enumerate(rows, start=1):
        item["payment_id"] = f"RCPT-WAU-{run_stamp}-{sequence:04d}"
    return rows


def write_output_from_template(
    benchmark_path: Path,
    column_order: list[str],
    rows: list[dict[str, object]],
    output_path: Path,
) -> None:
    workbook = load_workbook(benchmark_path)
    worksheet = workbook[workbook.sheetnames[0]]

    if worksheet.max_row > 1:
        worksheet.delete_rows(2, worksheet.max_row - 1)

    for row in rows:
        worksheet.append([row.get(column) for column in column_order])

    output_path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output_path)


def main() -> int:
    args = parse_args()

    wau_path = args.wau.expanduser().resolve()
    benchmark_path = args.benchmark.expanduser().resolve()
    orgunit_path = args.orgunit.expanduser().resolve()
    output_dir = args.output_dir.expanduser().resolve()

    for path in (wau_path, benchmark_path, orgunit_path):
        if not path.exists():
            raise FileNotFoundError(f"File not found: {path}")

    benchmark, fsp_name, column_order = load_benchmark_metadata(benchmark_path)
    entitlement_map = build_entitlement_map(benchmark)
    admin_lookup = build_admin_lookup(orgunit_path)
    roster = prepare_wau_roster(wau_path)
    selected_workers, selection_stats = select_workers(
        roster=roster,
        skip_count=max(args.skip_count, 0),
        extra_count=max(args.extra_count, 0),
        seed=args.seed,
    )

    missing_titles = sorted(
        title
        for title in selected_workers["title"].dropna().map(clean_text).unique().tolist()
        if title not in entitlement_map
    )
    if missing_titles:
        raise DataQualityError(
            "Selected Wau workers include titles missing from the benchmark entitlement map: "
            + ", ".join(missing_titles)
        )

    now = datetime.now()
    run_stamp = now.strftime("%d-%m-%Y-%H%M%S")
    output_path = output_dir / f"wau_payment_plan.{run_stamp}.xlsx"

    rows = build_output_rows(
        workers=selected_workers,
        entitlement_map=entitlement_map,
        admin_lookup=admin_lookup,
        fsp_name=fsp_name,
        run_stamp=run_stamp,
    )
    write_output_from_template(
        benchmark_path=benchmark_path,
        column_order=column_order,
        rows=rows,
        output_path=output_path,
    )

    print(f"Created: {output_path}")
    for key, value in selection_stats.items():
        print(f"{key}: {value}")
    print(f"fsp_name: {fsp_name}")
    print(f"payment_rows_written: {len(rows)}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
