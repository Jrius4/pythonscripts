#!/usr/bin/env python3
from __future__ import annotations

import argparse
import math
import re
import unicodedata
import warnings
from copy import copy
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter


REQUIRED_ORGUNIT_COLUMNS = ("State", "County", "Payam")
REQUIRED_HSPTP_COLUMNS = (
    "State",
    "State_PCode",
    "County",
    "County_PCode",
    "Payam",
    "Payam_PCode",
)


class MappingConflictError(RuntimeError):
    """Raised when the source workbook contains conflicting non-empty PCodes."""


def parse_args() -> argparse.Namespace:
    base_dir = Path(__file__).resolve().parent
    input_dir = base_dir / "inputs"

    parser = argparse.ArgumentParser(
        description=(
            "Map State, County, and Payam PCodes from the HSPTP workbook "
            "into the OrgUnit workbook and write a new OrgUnit_pcode.xlsx file."
        )
    )
    parser.add_argument(
        "--orgunit",
        type=Path,
        default=input_dir / "OrgUnit.xlsx",
        help="Path to the OrgUnit workbook.",
    )
    parser.add_argument(
        "--hsptp",
        type=Path,
        default=input_dir / "HSPTP_dataset.xlsx",
        help="Path to the HSPTP workbook.",
    )
    parser.add_argument(
        "--output",
        type=Path,
        default=input_dir / "OrgUnit_pcode.xlsx",
        help="Path for the generated workbook.",
    )
    parser.add_argument(
        "--strict",
        action="store_true",
        help="Fail if any Payam row cannot be mapped exactly.",
    )
    return parser.parse_args()


def normalize_text(value: object) -> str | None:
    if value is None:
        return None
    if isinstance(value, float) and math.isnan(value):
        return None

    text = unicodedata.normalize("NFKC", str(value))
    text = text.replace("\xa0", " ")
    text = " ".join(text.split()).strip()
    return text.casefold() if text else None


def clean_output_value(value: object) -> str | None:
    if value is None:
        return None
    if isinstance(value, float) and math.isnan(value):
        return None

    text = unicodedata.normalize("NFKC", str(value))
    text = text.replace("\xa0", " ").strip()
    return text or None


def display_admin_name(value: object) -> str | None:
    text = clean_output_value(value)
    if not text:
        return None

    text = re.sub(r"\s+(State|County|Payam)$", "", text, flags=re.IGNORECASE).strip()
    return text or None


def build_header_map(worksheet) -> dict[str, int]:
    header_map: dict[str, int] = {}
    for column in range(1, worksheet.max_column + 1):
        value = worksheet.cell(row=1, column=column).value
        normalized = clean_output_value(value)
        if normalized:
            header_map[normalized] = column
    return header_map


def find_orgunit_sheet(workbook):
    matches = []
    for worksheet in workbook.worksheets:
        headers = build_header_map(worksheet)
        if all(column in headers for column in REQUIRED_ORGUNIT_COLUMNS):
            matches.append((worksheet, headers))

    if not matches:
        raise ValueError(
            "Could not find a worksheet in the OrgUnit workbook with headers: "
            f"{', '.join(REQUIRED_ORGUNIT_COLUMNS)}"
        )

    return matches[0]


def read_hsptp_sheets(path: Path) -> tuple[list[str], dict[str, str], dict[tuple[str, str], str], dict[tuple[str, str, str], str]]:
    warnings.filterwarnings(
        "ignore",
        message="Slicer List extension is not supported and will be removed",
    )
    workbook = pd.read_excel(path, sheet_name=None, dtype=object)

    state_map: dict[str, str] = {}
    county_map: dict[tuple[str, str], str] = {}
    payam_map: dict[tuple[str, str, str], str] = {}
    sources: dict[str, dict[object, tuple[str, str, int]]] = {
        "state": {},
        "county": {},
        "payam": {},
    }
    used_sheets: list[str] = []

    def register(
        mapping_kind: str,
        mapping: dict,
        key: object,
        code_value: object,
        *,
        sheet_name: str,
        row_number: int,
    ) -> None:
        code = clean_output_value(code_value)
        if key is None or code is None:
            return

        existing = mapping.get(key)
        if existing is None:
            mapping[key] = code
            sources[mapping_kind][key] = (code, sheet_name, row_number)
            return

        if existing != code:
            previous_code, previous_sheet, previous_row = sources[mapping_kind][key]
            raise MappingConflictError(
                f"Conflicting {mapping_kind} mapping for {key!r}: "
                f"{previous_code!r} ({previous_sheet} row {previous_row}) vs "
                f"{code!r} ({sheet_name} row {row_number})"
            )

    for sheet_name, dataframe in workbook.items():
        if not all(column in dataframe.columns for column in REQUIRED_HSPTP_COLUMNS):
            continue

        used_sheets.append(sheet_name)

        for row_offset, row in enumerate(
            dataframe[list(REQUIRED_HSPTP_COLUMNS)].itertuples(index=False),
            start=2,
        ):
            state, state_pcode, county, county_pcode, payam, payam_pcode = row

            state_key = normalize_text(state)
            county_key = normalize_text(county)
            payam_key = normalize_text(payam)

            register(
                "state",
                state_map,
                state_key,
                state_pcode,
                sheet_name=sheet_name,
                row_number=row_offset,
            )
            register(
                "county",
                county_map,
                (state_key, county_key) if state_key and county_key else None,
                county_pcode,
                sheet_name=sheet_name,
                row_number=row_offset,
            )
            register(
                "payam",
                payam_map,
                (state_key, county_key, payam_key)
                if state_key and county_key and payam_key
                else None,
                payam_pcode,
                sheet_name=sheet_name,
                row_number=row_offset,
            )

    if not used_sheets:
        raise ValueError(
            "No HSPTP sheets were found with the required columns: "
            f"{', '.join(REQUIRED_HSPTP_COLUMNS)}"
        )

    return used_sheets, state_map, county_map, payam_map


def copy_cell_style(source_cell, target_cell) -> None:
    if source_cell.has_style:
        target_cell._style = copy(source_cell._style)
    if source_cell.number_format:
        target_cell.number_format = source_cell.number_format
    if source_cell.protection:
        target_cell.protection = copy(source_cell.protection)
    if source_cell.alignment:
        target_cell.alignment = copy(source_cell.alignment)


def insert_pcode_columns(worksheet, header_map: dict[str, int]) -> None:
    if any(column in header_map for column in ("State_PCode", "County_PCode", "Payam_PCode")):
        raise ValueError("The OrgUnit sheet already contains one or more *_PCode columns.")

    insertion_plan = [
        ("Payam_PCode", header_map["Payam"] + 1, header_map["Payam"]),
        ("County_PCode", header_map["County"] + 1, header_map["County"]),
        ("State_PCode", header_map["State"] + 1, header_map["State"]),
    ]

    for new_header, insert_at, template_column in insertion_plan:
        worksheet.insert_cols(insert_at, amount=1)

        source_letter = get_column_letter(template_column)
        target_letter = get_column_letter(insert_at)
        source_width = worksheet.column_dimensions[source_letter].width
        if source_width is not None:
            worksheet.column_dimensions[target_letter].width = source_width

        for row_number in range(1, worksheet.max_row + 1):
            source_cell = worksheet.cell(row=row_number, column=template_column)
            target_cell = worksheet.cell(row=row_number, column=insert_at)
            copy_cell_style(source_cell, target_cell)

        worksheet.cell(row=1, column=insert_at, value=new_header)


def apply_mappings(
    worksheet,
    state_map: dict[str, str],
    county_map: dict[tuple[str, str], str],
    payam_map: dict[tuple[str, str, str], str],
) -> dict[str, object]:
    header_map = build_header_map(worksheet)
    required_columns = (
        "State",
        "State_PCode",
        "County",
        "County_PCode",
        "Payam",
        "Payam_PCode",
    )
    missing = [column for column in required_columns if column not in header_map]
    if missing:
        raise ValueError(f"Missing required OrgUnit output columns: {', '.join(missing)}")

    stats = {
        "rows": 0,
        "state_matched": 0,
        "county_matched": 0,
        "payam_matched": 0,
        "unmatched_payams": [],
    }

    for row_number in range(2, worksheet.max_row + 1):
        stats["rows"] += 1

        state_cell = worksheet.cell(row=row_number, column=header_map["State"])
        county_cell = worksheet.cell(row=row_number, column=header_map["County"])
        payam_cell = worksheet.cell(row=row_number, column=header_map["Payam"])

        state_value = state_cell.value
        county_value = county_cell.value
        payam_value = payam_cell.value

        state_key = normalize_text(state_value)
        county_key = normalize_text(county_value)
        payam_key = normalize_text(payam_value)

        state_code = state_map.get(state_key)
        county_code = county_map.get((state_key, county_key))
        payam_code = payam_map.get((state_key, county_key, payam_key))

        worksheet.cell(row=row_number, column=header_map["State_PCode"]).value = state_code
        worksheet.cell(row=row_number, column=header_map["County_PCode"]).value = county_code
        worksheet.cell(row=row_number, column=header_map["Payam_PCode"]).value = payam_code

        state_cell.value = display_admin_name(state_value)
        county_cell.value = display_admin_name(county_value)
        payam_cell.value = display_admin_name(payam_value)

        if state_code:
            stats["state_matched"] += 1
        if county_code:
            stats["county_matched"] += 1
        if payam_code:
            stats["payam_matched"] += 1
        else:
            stats["unmatched_payams"].append(
                {
                    "row": row_number,
                    "State": clean_output_value(state_value),
                    "County": clean_output_value(county_value),
                    "Payam": clean_output_value(payam_value),
                }
            )

    return stats


def main() -> int:
    args = parse_args()

    orgunit_path = args.orgunit.expanduser().resolve()
    hsptp_path = args.hsptp.expanduser().resolve()
    output_path = args.output.expanduser().resolve()

    if not orgunit_path.exists():
        raise FileNotFoundError(f"OrgUnit workbook not found: {orgunit_path}")
    if not hsptp_path.exists():
        raise FileNotFoundError(f"HSPTP workbook not found: {hsptp_path}")

    used_sheets, state_map, county_map, payam_map = read_hsptp_sheets(hsptp_path)

    workbook = load_workbook(orgunit_path)
    worksheet, header_map = find_orgunit_sheet(workbook)
    insert_pcode_columns(worksheet, header_map)
    stats = apply_mappings(worksheet, state_map, county_map, payam_map)

    output_path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output_path)

    print(f"Created: {output_path}")
    print(f"HSPTP sheets used: {', '.join(used_sheets)}")
    print(
        f"Mapped rows: {stats['rows']} | "
        f"State_PCode: {stats['state_matched']} | "
        f"County_PCode: {stats['county_matched']} | "
        f"Payam_PCode: {stats['payam_matched']}"
    )

    unmatched_payams = stats["unmatched_payams"]
    if unmatched_payams:
        print(f"Unmatched Payam rows: {len(unmatched_payams)}")
        for item in unmatched_payams[:10]:
            print(
                f"  Row {item['row']}: "
                f"{item['State']} | {item['County']} | {item['Payam']}"
            )
        if args.strict:
            raise RuntimeError("Strict mode enabled and one or more Payam rows were unmatched.")
    else:
        print("Unmatched Payam rows: 0")

    return 0


if __name__ == "__main__":
    raise SystemExit(main())
